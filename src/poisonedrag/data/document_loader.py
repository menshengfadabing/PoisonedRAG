"""
通用文档加载器模块

支持多种文件格式的文本提取：
- PDF (.pdf) — pypdf
- Word (.docx) — python-docx
- PowerPoint (.pptx) — python-pptx
- Markdown (.md) — 纯文本读取
- JSON (.json) — 结构化解析
- 纯文本 (.txt) — 纯文本读取
"""

import json
import os
from typing import List, Optional, Dict, Any
from pathlib import Path

from langchain_core.documents import Document


# ============================================================
# 支持的文件扩展名映射
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".md": "markdown",
    ".json": "json",
    ".txt": "text",
}


def load_pdf(file_path: str) -> List[Document]:
    """从 PDF 文件提取文本"""
    from pypdf import PdfReader

    documents = []
    reader = PdfReader(file_path)
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            documents.append(Document(
                page_content=text,
                metadata={
                    "source": file_path,
                    "type": "knowledge",
                    "page": i + 1,
                }
            ))
    return documents


def load_docx(file_path: str) -> List[Document]:
    """从 Word (.docx) 文件提取文本"""
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    paragraphs = []
    current_section = ""

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            if current_section.strip():
                paragraphs.append(current_section.strip())
                current_section = ""
            continue
        # 按标题样式分段
        if para.style.name.startswith("Heading"):
            if current_section.strip():
                paragraphs.append(current_section.strip())
            current_section = text + "\n\n"
        else:
            current_section += text + "\n"

    if current_section.strip():
        paragraphs.append(current_section.strip())

    documents = []
    for i, section in enumerate(paragraphs):
        documents.append(Document(
            page_content=section,
            metadata={
                "source": file_path,
                "type": "knowledge",
                "section": i + 1,
            }
        ))
    return documents


def load_pptx(file_path: str) -> List[Document]:
    """从 PowerPoint (.pptx) 文件提取文本"""
    from pptx import Presentation

    prs = Presentation(file_path)
    documents = []

    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        content = "\n".join(texts)
        if content.strip():
            documents.append(Document(
                page_content=content,
                metadata={
                    "source": file_path,
                    "type": "knowledge",
                    "slide": i + 1,
                }
            ))
    return documents


def load_markdown(file_path: str) -> List[Document]:
    """从 Markdown 文件按标题分割"""
    import re

    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    sections = re.split(r'\n(?=#{1,2}\s)', content)
    documents = []

    for section in sections:
        section = section.strip()
        if not section:
            continue

        title_match = re.match(r'^(#{1,2})\s+(.+)$', section)
        title = title_match.group(2).strip() if title_match else "未命名"

        documents.append(Document(
            page_content=section,
            metadata={
                "source": file_path,
                "title": title,
                "type": "knowledge",
            }
        ))
    return documents


def load_json(file_path: str) -> List[Document]:
    """从 JSON 文件加载文档（支持数组和单对象）"""
    documents = []

    with open(file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 支持单对象和数组两种格式
    if isinstance(data, dict):
        data = [data]

    for item in data:
        content = item.get("content", "")
        metadata = item.get("metadata", {})
        metadata["source"] = metadata.get("source", file_path)
        metadata["type"] = "knowledge"

        doc = Document(page_content=content, metadata=metadata)
        documents.append(doc)

    return documents


def load_text(file_path: str, chunk_size: int = 500) -> List[Document]:
    """从纯文本文件按段落/块大小分割"""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    documents = []
    paragraphs = content.split('\n\n')

    current_chunk = ""
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue

        if len(current_chunk) + len(para) <= chunk_size:
            current_chunk += para + "\n\n"
        else:
            if current_chunk:
                documents.append(Document(
                    page_content=current_chunk.strip(),
                    metadata={"source": file_path, "type": "knowledge"},
                ))
            current_chunk = para + "\n\n"

    if current_chunk.strip():
        documents.append(Document(
            page_content=current_chunk.strip(),
            metadata={"source": file_path, "type": "knowledge"},
        ))
    return documents


def load_file(file_path: str, chunk_size: int = 500) -> List[Document]:
    """
    自动检测文件类型并加载文档

    Args:
        file_path: 文件路径
        chunk_size: 文本文件分块大小

    Returns:
        Document 列表

    Raises:
        ValueError: 不支持的文件格式
    """
    ext = Path(file_path).suffix.lower()
    file_type = SUPPORTED_EXTENSIONS.get(ext)

    if file_type is None:
        supported = ", ".join(SUPPORTED_EXTENSIONS.keys())
        raise ValueError(
            f"不支持的文件格式: {ext}，支持的格式: {supported}"
        )

    loaders = {
        "pdf": load_pdf,
        "docx": load_docx,
        "pptx": load_pptx,
        "markdown": load_markdown,
        "json": load_json,
        "text": lambda p: load_text(p, chunk_size),
    }

    return loaders[file_type](file_path)


def load_directory(
    directory: str,
    recursive: bool = True,
    chunk_size: int = 500,
    extensions: Optional[List[str]] = None,
) -> List[Document]:
    """
    从目录加载所有支持的文档

    Args:
        directory: 目录路径
        recursive: 是否递归
        chunk_size: 文本文件分块大小
        extensions: 允许的文件扩展名列表，None 表示全部支持

    Returns:
        Document 列表
    """
    documents = []

    if not os.path.exists(directory):
        return documents

    for root, dirs, files in os.walk(directory):
        if not recursive and root != directory:
            continue

        for file in files:
            file_path = os.path.join(root, file)
            ext = Path(file).suffix.lower()

            if extensions and ext not in extensions:
                continue
            if ext not in SUPPORTED_EXTENSIONS:
                continue

            try:
                docs = load_file(file_path, chunk_size)
                documents.extend(docs)
            except Exception as e:
                print(f"加载文件 {file_path} 失败: {e}")

    return documents
